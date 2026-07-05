import sys
import tempfile
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))


def check_file_exists(label, path):
    if path.exists():
        print(f"[OK] {label} exists")
        return 0

    print(f"[NG] {label} missing -> {path}")
    return 1


def check_true(label, condition):
    if condition:
        print(f"[OK] {label}")
        return 0

    print(f"[NG] {label}")
    return 1


def check_words(label, path, words):
    if not path.exists():
        print(f"[NG] {label} missing -> {path}")
        return 1

    text = path.read_text(encoding="utf-8")
    missing_words = [word for word in words if word not in text]

    if missing_words:
        print(f"[NG] {label} missing words -> {missing_words}")
        return 1

    print(f"[OK] {label} required words")
    return 0


def get_db_timestamp():
    db_path = PROJECT_ROOT / "data" / "personal_ai.db"
    return db_path.stat().st_mtime_ns if db_path.exists() else None


def check_db_unchanged(before):
    after = get_db_timestamp()
    return check_true("check_ver2410 does not write DB", before == after)


def check_check_script_has_no_forbidden_words():
    source = (PROJECT_ROOT / "scripts" / "check_ver2410.py").read_text(encoding="utf-8")
    forbidden_words = [
        "req" + "uests",
        "url" + "lib",
        "url" + "open",
        "ask_" + "llm(",
        "Base.metadata." + "create_all",
        "create_" + "database(",
        "reset --" + "hard",
        "force " + "push",
        "re" + "base",
    ]
    found_words = [word for word in forbidden_words if word in source]

    if found_words:
        print(f"[NG] check_ver2410 has forbidden words -> {found_words}")
        return 1

    print("[OK] check_ver2410 has no forbidden words")
    return 0


def check_imports():
    try:
        from core.pptx_reader_adapter import PPTXReaderAdapter, extract_pptx_path_from_text
        from core.pptx_reader_confirmation import (
            build_pptx_reader_confirmation,
            resolve_pptx_reader_confirmation_response,
        )
        from core.confirmation_intent import resolve_confirmation_intent
        from core.personal_ai import PersonalAI
        from core.tool_adapter import ToolRequest, ToolResult
        from core.tool_runner import run_tool_request
    except Exception as error:
        print(f"[NG] imports failed -> {error}")
        return 1

    PPTXReaderAdapter
    extract_pptx_path_from_text
    build_pptx_reader_confirmation
    resolve_pptx_reader_confirmation_response
    resolve_confirmation_intent
    PersonalAI
    ToolRequest
    ToolResult
    run_tool_request

    print("[OK] imports")
    return 0


def check_path_extraction():
    from core.pptx_reader_adapter import extract_pptx_path_from_text

    failures = 0
    failures += check_true(
        "bare pptx path detected",
        extract_pptx_path_from_text("このPowerPointを読んで docs/report.pptx")
        == "docs/report.pptx",
    )
    failures += check_true(
        "space-containing windows path detected",
        extract_pptx_path_from_text(r"C:\Users\me\Downloads\pptx_report (1).pptx")
        == r"C:\Users\me\Downloads\pptx_report (1).pptx",
    )
    failures += check_true(
        "no path returns None",
        extract_pptx_path_from_text("このPowerPointを読んで") is None,
    )
    failures += check_true(
        "non pptx extension is ignored",
        extract_pptx_path_from_text("note.txt を読んで") is None,
    )

    return failures


def _build_pptx(tmp_path, slides_text, include_table=False, include_notes=False):
    from pptx import Presentation

    presentation = Presentation()
    layout = presentation.slide_layouts[1]

    for title, body_text in slides_text:
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text_frame.text = body_text

        if include_notes:
            slide.notes_slide.notes_text_frame.text = "テスト用ノートです。"

    if include_table:
        blank_layout = presentation.slide_layouts[6]
        table_slide = presentation.slides.add_slide(blank_layout)
        table_shape = table_slide.shapes.add_table(1, 2, 0, 0, 2000000, 500000)
        table_shape.table.cell(0, 0).text = "Item"
        table_shape.table.cell(0, 1).text = "Value"

    path = tmp_path / "note.pptx"
    presentation.save(str(path))
    return path


def check_pptx_reader_adapter(tmp_path):
    from core.pptx_reader_adapter import PPTXReaderAdapter
    from core.tool_adapter import ToolRequest, ToolResult

    good_file = _build_pptx(
        tmp_path,
        [("Title", "PowerPoint Readerの動作確認用テキストです。")],
        include_table=True,
        include_notes=True,
    )

    adapter = PPTXReaderAdapter()
    failures = 0

    success_request = ToolRequest(
        tool_name="pptx_reader",
        action="read_pptx",
        input_data={"path": str(good_file)},
        requires_permission=True,
        risk_level="medium",
    )
    success_result = adapter.execute(success_request)
    failures += check_true("adapter is available", adapter.is_available() is True)
    failures += check_true("execute returns ToolResult", isinstance(success_result, ToolResult))
    failures += check_true("success result success", success_result.success is True)
    failures += check_true("success result executed", success_result.executed is True)
    failures += check_true(
        "success result has file_name",
        success_result.data.get("file_name") == "note.pptx",
    )
    failures += check_true(
        "success result has text preview",
        "動作確認" in success_result.data.get("text_preview", ""),
    )
    failures += check_true(
        "success result includes table text",
        "Item | Value" in success_result.data.get("text", ""),
    )
    failures += check_true(
        "success result includes notes text",
        "[ノート]" in success_result.data.get("text", ""),
    )
    failures += check_true(
        "success result has slide_count",
        success_result.data.get("slide_count") == 2,
    )

    missing_request = ToolRequest(
        tool_name="pptx_reader",
        action="read_pptx",
        input_data={"path": str(tmp_path / "does_not_exist.pptx")},
        requires_permission=True,
        risk_level="medium",
    )
    missing_result = adapter.execute(missing_request)
    failures += check_true("missing file fails", missing_result.success is False)
    failures += check_true("missing file not executed", missing_result.executed is False)

    wrong_extension_request = ToolRequest(
        tool_name="pptx_reader",
        action="read_pptx",
        input_data={"path": str(tmp_path / "note.txt")},
        requires_permission=True,
        risk_level="medium",
    )
    wrong_extension_result = adapter.execute(wrong_extension_request)
    failures += check_true("wrong extension fails", wrong_extension_result.success is False)
    failures += check_true("wrong extension not executed", wrong_extension_result.executed is False)

    corrupt_file = tmp_path / "corrupt.pptx"
    corrupt_file.write_bytes(b"not a real pptx file")
    corrupt_request = ToolRequest(
        tool_name="pptx_reader",
        action="read_pptx",
        input_data={"path": str(corrupt_file)},
        requires_permission=True,
        risk_level="medium",
    )
    corrupt_result = adapter.execute(corrupt_request)
    failures += check_true("corrupt pptx fails honestly", corrupt_result.success is False)
    failures += check_true("corrupt pptx executed True", corrupt_result.executed is True)

    return failures


def check_confirmation_flow(tmp_path):
    from core.pptx_reader_confirmation import (
        build_pptx_reader_confirmation,
        resolve_pptx_reader_confirmation_response,
    )

    target_file = _build_pptx(tmp_path, [("Title", "確認フローのテスト本文です。")])

    confirmation = build_pptx_reader_confirmation(f"このPowerPointを読んで {target_file}")
    failures = 0
    failures += check_true("confirmation needs_confirmation", confirmation.needs_confirmation is True)
    failures += check_true("confirmation has path", confirmation.path == str(target_file))
    failures += check_true("confirmation tool_name", confirmation.tool_name == "pptx_reader")

    pending = confirmation.to_dict()
    pending["original_input"] = f"このPowerPointを読んで {target_file}"

    confirmed = resolve_pptx_reader_confirmation_response("いいよ", pending)
    failures += check_true("confirmed status", confirmed.status == "confirmed")
    failures += check_true(
        "confirmed candidate tool_name",
        confirmed.tool_request_candidate["tool_name"] == "pptx_reader",
    )

    cancelled = resolve_pptx_reader_confirmation_response("キャンセル", pending)
    failures += check_true("cancelled status", cancelled.status == "cancelled")

    unclear = resolve_pptx_reader_confirmation_response("よくわからない", pending)
    failures += check_true("unclear status", unclear.status == "unclear")

    no_pending = resolve_pptx_reader_confirmation_response("いいよ", None)
    failures += check_true("no pending status", no_pending.status == "no_pending")

    return failures


def check_tool_runner_permission(tmp_path):
    from core.tool_adapter import ToolRequest
    from core.tool_execution_log import ToolExecutionLogger
    from core.tool_registry import ToolRegistry
    from core.tool_runner import run_tool_request
    from core.pptx_reader_adapter import PPTXReaderAdapter

    target_file = _build_pptx(tmp_path, [("Title", "Permission経由の実行確認です。")])

    registry = ToolRegistry()
    registry.register(PPTXReaderAdapter())
    logger = ToolExecutionLogger()
    request = ToolRequest(
        tool_name="pptx_reader",
        action="read_pptx",
        input_data={"path": str(target_file)},
        requires_permission=True,
        risk_level="medium",
    )

    blocked = run_tool_request(
        request,
        registry=registry,
        user_confirmed=False,
        logger=logger,
    )
    executed = run_tool_request(
        request,
        registry=registry,
        user_confirmed=True,
        logger=logger,
    )

    failures = 0
    failures += check_true("permission needed", blocked.permission_decision.needs_user_confirmation is True)
    failures += check_true("no permission not executed", blocked.result.executed is False)
    failures += check_true("confirmed executed", executed.result.executed is True)
    failures += check_true("confirmed success", executed.result.success is True)
    failures += check_true("logger recorded both results", logger.count() == 2)

    return failures


def fake_llm_responder(messages):
    system_texts = [message["content"] for message in messages if message["role"] == "system"]
    user_texts = [message["content"] for message in messages if message["role"] == "user"]
    document_in_context = any("PPTX_FAKE_MARKER" in text for text in system_texts)
    last_user_text = user_texts[-1] if user_texts else ""
    return f"FAKE_LLM_ANSWER|document_in_context={document_in_context}|user={last_user_text}"


def build_lightweight_personal_ai():
    from core.capability_manager import CapabilityManager
    from core.input_analyzer import InputAnalyzer
    from core.personal_ai import PersonalAI
    from core.tool_execution_log import ToolExecutionLogger

    ai = PersonalAI.__new__(PersonalAI)
    ai.input_analyzer = InputAnalyzer()
    ai.capability_manager = CapabilityManager()
    ai.web_reader_fetcher = None
    ai.ask_llm_func = fake_llm_responder
    ai.tool_registry = ai.build_tool_registry(None)
    ai.tool_execution_logger = ToolExecutionLogger()
    ai.pending_preference_confirmation = None
    ai.pending_memory_candidate_confirmation = None
    ai.pending_memory_candidate_edit = None
    ai.pending_web_reader_confirmation = None
    ai.pending_file_reader_confirmation = None
    ai.pending_pdf_reader_confirmation = None
    ai.pending_word_reader_confirmation = None
    ai.pending_excel_reader_confirmation = None
    ai.pending_pptx_reader_confirmation = None
    ai.last_web_reader_confirmation = None
    ai.last_web_reader_confirmation_result = None
    ai.last_web_reader_tool_request_candidate = None
    ai.last_web_reader_tool_run_result = None
    ai.last_web_reader_tool_result = None
    ai.last_file_reader_confirmation = None
    ai.last_file_reader_confirmation_result = None
    ai.last_file_reader_tool_request_candidate = None
    ai.last_file_reader_tool_run_result = None
    ai.last_file_reader_tool_result = None
    ai.last_pdf_reader_confirmation = None
    ai.last_pdf_reader_confirmation_result = None
    ai.last_pdf_reader_tool_request_candidate = None
    ai.last_pdf_reader_tool_run_result = None
    ai.last_pdf_reader_tool_result = None
    ai.last_word_reader_confirmation = None
    ai.last_word_reader_confirmation_result = None
    ai.last_word_reader_tool_request_candidate = None
    ai.last_word_reader_tool_run_result = None
    ai.last_word_reader_tool_result = None
    ai.last_excel_reader_confirmation = None
    ai.last_excel_reader_confirmation_result = None
    ai.last_excel_reader_tool_request_candidate = None
    ai.last_excel_reader_tool_run_result = None
    ai.last_excel_reader_tool_result = None
    ai.last_pptx_reader_confirmation = None
    ai.last_pptx_reader_confirmation_result = None
    ai.last_pptx_reader_tool_request_candidate = None
    ai.last_pptx_reader_tool_run_result = None
    ai.last_pptx_reader_tool_result = None
    ai.last_read_document = None
    ai.last_tool_execution_record = None
    ai.last_memory_candidate = None
    ai.last_memory_candidate_confirmation_result = None
    ai.last_memory_candidate_save_result = None
    ai.last_saved_memory_candidate_result = None
    ai.last_memory_usage = "normal"
    ai.last_conversation_mode = "chat"
    ai.last_conversation_quality_guidance = None
    ai.last_input_analysis = None
    ai.last_response = None
    ai.save_user_message = lambda _message: None
    ai.save_assistant_message = lambda _message: None
    return ai


def check_personal_ai_flow(tmp_path):
    # ファイル名に"powerpoint"を含めても問題ないことも合わせて確認します。
    target_file = _build_pptx(tmp_path, [("Title", "PPTX_FAKE_MARKERを含む本文です。結合テストです。")])
    renamed_file = target_file.rename(tmp_path / "my_powerpoint_slides (1).pptx")

    failures = 0

    ai_confirm = build_lightweight_personal_ai()
    confirmation_response = ai_confirm.handle_user_input(f"このPowerPointを読んで {renamed_file}")
    read_response = ai_confirm.handle_user_input("いいよ")
    failures += check_true(
        "confirmation shown for space+powerpoint-containing filename",
        str(renamed_file) in confirmation_response,
    )
    failures += check_true("reads with PowerPoint Reader", "PowerPoint Readerで読み取りました" in read_response)
    failures += check_true("shows file content", "結合テスト" in read_response)
    failures += check_true("pending cleared", ai_confirm.pending_pptx_reader_confirmation is None)
    failures += check_true(
        "last_read_document stored as pptx_reader",
        ai_confirm.last_read_document is not None
        and ai_confirm.last_read_document.get("source_tool") == "pptx_reader",
    )

    ai_summarize = build_lightweight_personal_ai()
    ai_summarize.handle_user_input(f"{renamed_file} 内容を要約してください")
    summarize_response = ai_summarize.handle_user_input("はい")
    failures += check_true("combined flow calls fake LLM", "FAKE_LLM_ANSWER" in summarize_response)
    failures += check_true(
        "combined flow LLM sees document",
        "document_in_context=True" in summarize_response,
    )

    ai_cancel = build_lightweight_personal_ai()
    ai_cancel.handle_user_input(f"このPowerPointを読んで {renamed_file}")
    cancel_response = ai_cancel.handle_user_input("キャンセル")
    failures += check_true("cancel not executed", "PowerPoint Readerは実行しませんでした" in cancel_response)
    failures += check_true("cancel pending cleared", ai_cancel.pending_pptx_reader_confirmation is None)

    ai_missing = build_lightweight_personal_ai()
    missing_path = tmp_path / "missing.pptx"
    ai_missing.handle_user_input(f"このPowerPointを読んで {missing_path}")
    missing_response = ai_missing.handle_user_input("はい")
    failures += check_true(
        "missing file reports not executed",
        "PowerPoint Readerは実行されませんでした" in missing_response,
    )

    return failures


def main():
    failures = 0
    db_timestamp_before = get_db_timestamp()

    docs = PROJECT_ROOT / "docs"
    scripts = PROJECT_ROOT / "scripts"
    core = PROJECT_ROOT / "core"

    files = [
        ("pptx_reader_adapter.py", core / "pptx_reader_adapter.py"),
        ("pptx_reader_confirmation.py", core / "pptx_reader_confirmation.py"),
        ("confirmation_intent.py", core / "confirmation_intent.py"),
        ("personal_ai.py", core / "personal_ai.py"),
        ("check_ver2410.py", scripts / "check_ver2410.py"),
        ("run_ver2_checks.py", scripts / "run_ver2_checks.py"),
        ("requirements.txt", PROJECT_ROOT / "requirements.txt"),
        ("changelog.md", docs / "changelog.md"),
        ("project_handoff.md", docs / "project_handoff.md"),
    ]

    for label, path in files:
        failures += check_file_exists(label, path)

    failures += check_imports()
    failures += check_path_extraction()

    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_path = Path(tmp_dir)
        failures += check_pptx_reader_adapter(tmp_path)
        failures += check_confirmation_flow(tmp_path)
        failures += check_tool_runner_permission(tmp_path)
        failures += check_personal_ai_flow(tmp_path)

    failures += check_words(
        "requirements.txt",
        PROJECT_ROOT / "requirements.txt",
        [
            "python-pptx",
        ],
    )

    failures += check_words(
        "run_ver2_checks.py",
        scripts / "run_ver2_checks.py",
        [
            "check_ver2410.py",
        ],
    )

    failures += check_words(
        "changelog.md",
        docs / "changelog.md",
        [
            "## Version 2.41",
            "PowerPoint Reader",
            "scripts/check_ver2410.py",
        ],
    )

    failures += check_words(
        "project_handoff.md",
        docs / "project_handoff.md",
        [
            "現在バージョン:",
            "Ver2.41: PowerPoint Reader最小実装",
        ],
    )

    failures += check_check_script_has_no_forbidden_words()
    failures += check_db_unchanged(db_timestamp_before)

    if failures:
        print(f"{failures} checks failed.")
        raise SystemExit(1)

    print("All checks passed.")


if __name__ == "__main__":
    main()
