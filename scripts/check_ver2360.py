import sys
import tempfile
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))


def check_file_exists(label, path):
    if path.exists():
        print(f"[OK] {label} exists")
        return 0

    print(f"[NG] {label} missing -> {path}")
    return 1


def check_true(label, condition):
    if condition:
        print(f"[OK] {label}")
        return 0

    print(f"[NG] {label}")
    return 1


def check_words(label, path, words):
    if not path.exists():
        print(f"[NG] {label} missing -> {path}")
        return 1

    text = path.read_text(encoding="utf-8")
    missing_words = [word for word in words if word not in text]

    if missing_words:
        print(f"[NG] {label} missing words -> {missing_words}")
        return 1

    print(f"[OK] {label} required words")
    return 0


def get_db_timestamp():
    db_path = PROJECT_ROOT / "data" / "personal_ai.db"
    return db_path.stat().st_mtime_ns if db_path.exists() else None


def check_db_unchanged(before):
    after = get_db_timestamp()
    return check_true("check_ver2360 does not write DB", before == after)


def check_check_script_has_no_forbidden_words():
    source = (PROJECT_ROOT / "scripts" / "check_ver2360.py").read_text(encoding="utf-8")
    forbidden_words = [
        "req" + "uests",
        "url" + "lib",
        "url" + "open",
        "ask_" + "llm(",
        "Base.metadata." + "create_all",
        "create_" + "database(",
        "reset --" + "hard",
        "force " + "push",
        "re" + "base",
    ]
    found_words = [word for word in forbidden_words if word in source]

    if found_words:
        print(f"[NG] check_ver2360 has forbidden words -> {found_words}")
        return 1

    print("[OK] check_ver2360 has no forbidden words")
    return 0


def check_imports():
    try:
        from core.word_reader_adapter import WordReaderAdapter, extract_word_path_from_text
        from core.word_reader_confirmation import (
            build_word_reader_confirmation,
            resolve_word_reader_confirmation_response,
        )
        from core.confirmation_intent import resolve_confirmation_intent
        from core.personal_ai import PersonalAI
        from core.tool_adapter import ToolRequest, ToolResult
        from core.tool_runner import run_tool_request
    except Exception as error:
        print(f"[NG] imports failed -> {error}")
        return 1

    WordReaderAdapter
    extract_word_path_from_text
    build_word_reader_confirmation
    resolve_word_reader_confirmation_response
    resolve_confirmation_intent
    PersonalAI
    ToolRequest
    ToolResult
    run_tool_request

    print("[OK] imports")
    return 0


def check_path_extraction():
    from core.word_reader_adapter import extract_word_path_from_text

    failures = 0
    failures += check_true(
        "bare docx path detected",
        extract_word_path_from_text("このWordファイルを読んで docs/report.docx")
        == "docs/report.docx",
    )
    failures += check_true(
        "space-containing windows path detected",
        extract_word_path_from_text(r"C:\Users\me\Downloads\word_test (1).docx")
        == r"C:\Users\me\Downloads\word_test (1).docx",
    )
    failures += check_true(
        "no path returns None",
        extract_word_path_from_text("このWordファイルを読んで") is None,
    )
    failures += check_true(
        "non docx extension is ignored",
        extract_word_path_from_text("note.txt を読んで") is None,
    )

    return failures


def _build_docx(tmp_path, paragraphs, include_table=False):
    from docx import Document

    document = Document()

    for paragraph in paragraphs:
        document.add_paragraph(paragraph)

    if include_table:
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Item"
        table.rows[0].cells[1].text = "Value"

    path = tmp_path / "note.docx"
    document.save(str(path))
    return path


def check_word_reader_adapter(tmp_path):
    from core.word_reader_adapter import WordReaderAdapter
    from core.tool_adapter import ToolRequest, ToolResult

    good_file = _build_docx(
        tmp_path,
        ["Word Readerの動作確認用テキストです。"],
        include_table=True,
    )

    adapter = WordReaderAdapter()
    failures = 0

    success_request = ToolRequest(
        tool_name="word_reader",
        action="read_word",
        input_data={"path": str(good_file)},
        requires_permission=True,
        risk_level="medium",
    )
    success_result = adapter.execute(success_request)
    failures += check_true("adapter is available", adapter.is_available() is True)
    failures += check_true("execute returns ToolResult", isinstance(success_result, ToolResult))
    failures += check_true("success result success", success_result.success is True)
    failures += check_true("success result executed", success_result.executed is True)
    failures += check_true(
        "success result has file_name",
        success_result.data.get("file_name") == "note.docx",
    )
    failures += check_true(
        "success result has text preview",
        "動作確認" in success_result.data.get("text_preview", ""),
    )
    failures += check_true(
        "success result includes table text",
        "Item | Value" in success_result.data.get("text", ""),
    )

    missing_request = ToolRequest(
        tool_name="word_reader",
        action="read_word",
        input_data={"path": str(tmp_path / "does_not_exist.docx")},
        requires_permission=True,
        risk_level="medium",
    )
    missing_result = adapter.execute(missing_request)
    failures += check_true("missing file fails", missing_result.success is False)
    failures += check_true("missing file not executed", missing_result.executed is False)

    wrong_extension_request = ToolRequest(
        tool_name="word_reader",
        action="read_word",
        input_data={"path": str(tmp_path / "note.txt")},
        requires_permission=True,
        risk_level="medium",
    )
    wrong_extension_result = adapter.execute(wrong_extension_request)
    failures += check_true("wrong extension fails", wrong_extension_result.success is False)
    failures += check_true("wrong extension not executed", wrong_extension_result.executed is False)

    corrupt_file = tmp_path / "corrupt.docx"
    corrupt_file.write_bytes(b"not a real docx file")
    corrupt_request = ToolRequest(
        tool_name="word_reader",
        action="read_word",
        input_data={"path": str(corrupt_file)},
        requires_permission=True,
        risk_level="medium",
    )
    corrupt_result = adapter.execute(corrupt_request)
    failures += check_true("corrupt docx fails honestly", corrupt_result.success is False)
    failures += check_true("corrupt docx executed True", corrupt_result.executed is True)

    return failures


def check_confirmation_flow(tmp_path):
    from core.word_reader_confirmation import (
        build_word_reader_confirmation,
        resolve_word_reader_confirmation_response,
    )

    target_file = _build_docx(tmp_path, ["確認フローのテスト本文です。"])

    confirmation = build_word_reader_confirmation(f"このWordファイルを読んで {target_file}")
    failures = 0
    failures += check_true("confirmation needs_confirmation", confirmation.needs_confirmation is True)
    failures += check_true("confirmation has path", confirmation.path == str(target_file))
    failures += check_true("confirmation tool_name", confirmation.tool_name == "word_reader")

    pending = confirmation.to_dict()
    pending["original_input"] = f"このWordファイルを読んで {target_file}"

    confirmed = resolve_word_reader_confirmation_response("いいよ", pending)
    failures += check_true("confirmed status", confirmed.status == "confirmed")
    failures += check_true(
        "confirmed candidate tool_name",
        confirmed.tool_request_candidate["tool_name"] == "word_reader",
    )

    cancelled = resolve_word_reader_confirmation_response("キャンセル", pending)
    failures += check_true("cancelled status", cancelled.status == "cancelled")

    unclear = resolve_word_reader_confirmation_response("よくわからない", pending)
    failures += check_true("unclear status", unclear.status == "unclear")

    no_pending = resolve_word_reader_confirmation_response("いいよ", None)
    failures += check_true("no pending status", no_pending.status == "no_pending")

    return failures


def check_tool_runner_permission(tmp_path):
    from core.tool_adapter import ToolRequest
    from core.tool_execution_log import ToolExecutionLogger
    from core.tool_registry import ToolRegistry
    from core.tool_runner import run_tool_request
    from core.word_reader_adapter import WordReaderAdapter

    target_file = _build_docx(tmp_path, ["Permission経由の実行確認です。"])

    registry = ToolRegistry()
    registry.register(WordReaderAdapter())
    logger = ToolExecutionLogger()
    request = ToolRequest(
        tool_name="word_reader",
        action="read_word",
        input_data={"path": str(target_file)},
        requires_permission=True,
        risk_level="medium",
    )

    blocked = run_tool_request(
        request,
        registry=registry,
        user_confirmed=False,
        logger=logger,
    )
    executed = run_tool_request(
        request,
        registry=registry,
        user_confirmed=True,
        logger=logger,
    )

    failures = 0
    failures += check_true("permission needed", blocked.permission_decision.needs_user_confirmation is True)
    failures += check_true("no permission not executed", blocked.result.executed is False)
    failures += check_true("confirmed executed", executed.result.executed is True)
    failures += check_true("confirmed success", executed.result.success is True)
    failures += check_true("logger recorded both results", logger.count() == 2)

    return failures


def fake_llm_responder(messages):
    system_texts = [message["content"] for message in messages if message["role"] == "system"]
    user_texts = [message["content"] for message in messages if message["role"] == "user"]
    document_in_context = any("WORD_FAKE_MARKER" in text for text in system_texts)
    last_user_text = user_texts[-1] if user_texts else ""
    return f"FAKE_LLM_ANSWER|document_in_context={document_in_context}|user={last_user_text}"


def build_lightweight_personal_ai():
    from core.capability_manager import CapabilityManager
    from core.input_analyzer import InputAnalyzer
    from core.personal_ai import PersonalAI
    from core.tool_execution_log import ToolExecutionLogger

    ai = PersonalAI.__new__(PersonalAI)
    ai.input_analyzer = InputAnalyzer()
    ai.capability_manager = CapabilityManager()
    ai.web_reader_fetcher = None
    ai.ask_llm_func = fake_llm_responder
    ai.tool_registry = ai.build_tool_registry(None)
    ai.tool_execution_logger = ToolExecutionLogger()
    ai.pending_preference_confirmation = None
    ai.pending_memory_candidate_confirmation = None
    ai.pending_memory_candidate_edit = None
    ai.pending_web_reader_confirmation = None
    ai.pending_file_reader_confirmation = None
    ai.pending_pdf_reader_confirmation = None
    ai.pending_word_reader_confirmation = None
    ai.last_web_reader_confirmation = None
    ai.last_web_reader_confirmation_result = None
    ai.last_web_reader_tool_request_candidate = None
    ai.last_web_reader_tool_run_result = None
    ai.last_web_reader_tool_result = None
    ai.last_file_reader_confirmation = None
    ai.last_file_reader_confirmation_result = None
    ai.last_file_reader_tool_request_candidate = None
    ai.last_file_reader_tool_run_result = None
    ai.last_file_reader_tool_result = None
    ai.last_pdf_reader_confirmation = None
    ai.last_pdf_reader_confirmation_result = None
    ai.last_pdf_reader_tool_request_candidate = None
    ai.last_pdf_reader_tool_run_result = None
    ai.last_pdf_reader_tool_result = None
    ai.last_word_reader_confirmation = None
    ai.last_word_reader_confirmation_result = None
    ai.last_word_reader_tool_request_candidate = None
    ai.last_word_reader_tool_run_result = None
    ai.last_word_reader_tool_result = None
    ai.last_read_document = None
    ai.last_tool_execution_record = None
    ai.last_memory_candidate = None
    ai.last_memory_candidate_confirmation_result = None
    ai.last_memory_candidate_save_result = None
    ai.last_saved_memory_candidate_result = None
    ai.last_memory_usage = "normal"
    ai.last_conversation_mode = "chat"
    ai.last_conversation_quality_guidance = None
    ai.last_input_analysis = None
    ai.last_response = None
    ai.save_user_message = lambda _message: None
    ai.save_assistant_message = lambda _message: None
    return ai


def check_personal_ai_flow(tmp_path):
    # ファイル名に"word"を含めて、文書作成判定との衝突がないことも合わせて確認します。
    target_file = _build_docx(tmp_path, ["WORD_FAKE_MARKERを含む本文です。結合テストです。"])
    word_named_file = target_file.rename(tmp_path / "word_report (1).docx")

    failures = 0

    ai_confirm = build_lightweight_personal_ai()
    confirmation_response = ai_confirm.handle_user_input(f"このWordファイルを読んで {word_named_file}")
    read_response = ai_confirm.handle_user_input("いいよ")
    failures += check_true(
        "confirmation shown for space+word-containing filename",
        str(word_named_file) in confirmation_response,
    )
    failures += check_true("reads with Word Reader", "Word Readerで読み取りました" in read_response)
    failures += check_true("shows file content", "結合テスト" in read_response)
    failures += check_true("pending cleared", ai_confirm.pending_word_reader_confirmation is None)
    failures += check_true(
        "last_read_document stored as word_reader",
        ai_confirm.last_read_document is not None
        and ai_confirm.last_read_document.get("source_tool") == "word_reader",
    )

    ai_summarize = build_lightweight_personal_ai()
    ai_summarize.handle_user_input(f"{word_named_file} 内容を要約してください")
    summarize_response = ai_summarize.handle_user_input("はい")
    failures += check_true("combined flow calls fake LLM", "FAKE_LLM_ANSWER" in summarize_response)
    failures += check_true(
        "combined flow LLM sees document despite 'word' in filename",
        "document_in_context=True" in summarize_response,
    )

    ai_cancel = build_lightweight_personal_ai()
    ai_cancel.handle_user_input(f"このWordファイルを読んで {word_named_file}")
    cancel_response = ai_cancel.handle_user_input("キャンセル")
    failures += check_true("cancel not executed", "Word Readerは実行しませんでした" in cancel_response)
    failures += check_true("cancel pending cleared", ai_cancel.pending_word_reader_confirmation is None)

    ai_missing = build_lightweight_personal_ai()
    missing_path = tmp_path / "missing.docx"
    ai_missing.handle_user_input(f"このWordファイルを読んで {missing_path}")
    missing_response = ai_missing.handle_user_input("はい")
    failures += check_true(
        "missing file reports not executed",
        "Word Readerは実行されませんでした" in missing_response,
    )

    return failures


def check_word_excel_keyword_collision_fixed():
    from core.input_analyzer import InputAnalyzer
    from core.capability_manager import CapabilityManager

    ia = InputAnalyzer()
    cm = CapabilityManager()
    failures = 0

    filename_cases = [
        "word_report (1).docx を読んで要約して",
        "excel_summary.xlsx を読んで",
    ]

    for text in filename_cases:
        analysis = ia.analyze(text)
        capability = cm.analyze(text, analysis)
        failures += check_true(
            f"filename with word/excel substring not misclassified: {text}",
            analysis["requires_file_creation"] is False
            and capability["primary_capability"] != "document_creation_required",
        )

    failures += check_true(
        "real word document creation request still detected",
        ia.analyze("Wordで文書を作って")["requires_file_creation"] is True,
    )

    excel_text = "Excelで表を作って"
    excel_analysis = ia.analyze(excel_text)
    excel_capability = cm.analyze(excel_text, excel_analysis)
    failures += check_true(
        "real excel document creation request still detected",
        excel_capability["primary_capability"] == "document_creation_required",
    )

    return failures


def main():
    failures = 0
    db_timestamp_before = get_db_timestamp()

    docs = PROJECT_ROOT / "docs"
    scripts = PROJECT_ROOT / "scripts"
    core = PROJECT_ROOT / "core"

    files = [
        ("word_reader_adapter.py", core / "word_reader_adapter.py"),
        ("word_reader_confirmation.py", core / "word_reader_confirmation.py"),
        ("confirmation_intent.py", core / "confirmation_intent.py"),
        ("personal_ai.py", core / "personal_ai.py"),
        ("check_ver2360.py", scripts / "check_ver2360.py"),
        ("run_ver2_checks.py", scripts / "run_ver2_checks.py"),
        ("requirements.txt", PROJECT_ROOT / "requirements.txt"),
        ("changelog.md", docs / "changelog.md"),
        ("project_handoff.md", docs / "project_handoff.md"),
    ]

    for label, path in files:
        failures += check_file_exists(label, path)

    failures += check_imports()
    failures += check_path_extraction()
    failures += check_word_excel_keyword_collision_fixed()

    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_path = Path(tmp_dir)
        failures += check_word_reader_adapter(tmp_path)
        failures += check_confirmation_flow(tmp_path)
        failures += check_tool_runner_permission(tmp_path)
        failures += check_personal_ai_flow(tmp_path)

    failures += check_words(
        "requirements.txt",
        PROJECT_ROOT / "requirements.txt",
        [
            "python-docx",
        ],
    )

    failures += check_words(
        "run_ver2_checks.py",
        scripts / "run_ver2_checks.py",
        [
            "check_ver2360.py",
        ],
    )

    failures += check_words(
        "changelog.md",
        docs / "changelog.md",
        [
            "## Version 2.36",
            "Word Reader",
            "scripts/check_ver2360.py",
        ],
    )

    failures += check_words(
        "project_handoff.md",
        docs / "project_handoff.md",
        [
            "現在バージョン:",
            "Ver2.36: Word Reader最小実装",
        ],
    )

    failures += check_check_script_has_no_forbidden_words()
    failures += check_db_unchanged(db_timestamp_before)

    if failures:
        print(f"{failures} checks failed.")
        raise SystemExit(1)

    print("All checks passed.")


if __name__ == "__main__":
    main()
