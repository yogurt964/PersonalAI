from pathlib import Path

from core.file_path_extraction import extract_single_file_path_from_text
from core.tool_adapter import ToolAdapter, ToolResult


PPTX_EXTENSIONS = (".pptx",)


def extract_pptx_path_from_text(text):
    """テキストから、ユーザーが指定した単一のPowerPoint(.pptx)ファイルパスを取り出します。

    ディレクトリ全体やワイルドカードは対象にせず、拡張子で終わる1つのパスだけを拾います。
    スペースを含むファイル名（例: Windowsの重複ダウンロードによる「 (1).pptx」）にも対応します。
    """
    return extract_single_file_path_from_text(text, PPTX_EXTENSIONS)


class PPTXReaderAdapter(ToolAdapter):
    """最小実用版のPowerPoint Reader Adapterです。

    ユーザーが明示した単一の.pptxファイルだけを、確認後にテキスト抽出します。
    ディレクトリ探索、ワイルドカード、書き込み、実行は一切行いません。
    画像・図形内の装飾テキストや埋め込みオブジェクトの内容までは読み取らず、
    スライド上のテキスト（タイトル・本文・表・ノート）を対象にします。
    古い形式の.ppt（バイナリ）には対応しません。
    """

    def __init__(
        self,
        max_file_bytes=50 * 1024 * 1024,
        max_slides=100,
        max_text_chars=12000,
        presentation_reader_factory=None,
    ):
        super().__init__(
            name="pptx_reader",
            description="Read a single local PowerPoint (.pptx) file specified by the user.",
        )
        self.max_file_bytes = max_file_bytes
        self.max_slides = max_slides
        self.max_text_chars = max_text_chars
        # 自動checkでfakeのプレゼンテーション読み取りに差し替えられるようにします。
        self._presentation_reader_factory = presentation_reader_factory

    def is_available(self):
        return True

    def execute(self, request):
        """ToolRequestを受け取り、スライド本文をToolResultとして返します。"""
        path_text = self._get_path_from_request(request)

        if not path_text:
            return ToolResult(
                success=False,
                message="読み取るPowerPointファイルパスが見つからないため、PowerPoint Readerは実行していません。",
                data=None,
                error="missing_path",
                executed=False,
            )

        path = Path(path_text)

        if path.suffix.lower() != ".pptx":
            return ToolResult(
                success=False,
                message=(
                    "対応していない拡張子のため、PowerPoint Readerは実行していません。"
                    "（対応拡張子: .pptx。古い.ppt形式には対応していません）"
                ),
                data={"path": path_text},
                error="unsupported_extension",
                executed=False,
            )

        if not path.exists():
            return ToolResult(
                success=False,
                message="指定されたファイルが見つからないため、PowerPoint Readerは実行していません。",
                data={"path": path_text},
                error="file_not_found",
                executed=False,
            )

        if not path.is_file():
            return ToolResult(
                success=False,
                message="指定されたパスはファイルではないため、PowerPoint Readerは実行していません。",
                data={"path": path_text},
                error="not_a_file",
                executed=False,
            )

        try:
            size_bytes = path.stat().st_size
        except OSError as error:
            return ToolResult(
                success=False,
                message="ファイル情報の取得に失敗しました。読めたふりはしません。",
                data={"path": path_text},
                error=str(error),
                executed=True,
            )

        if size_bytes > self.max_file_bytes:
            return ToolResult(
                success=False,
                message=(
                    "ファイルサイズが上限を超えているため、PowerPoint Readerは実行していません。"
                    f"（上限: {self.max_file_bytes}バイト）"
                ),
                data={"path": path_text, "size_bytes": size_bytes},
                error="file_too_large",
                executed=False,
            )

        try:
            slide_count, text = self._extract_text(path)
        except Exception as error:
            return ToolResult(
                success=False,
                message="PowerPointファイルの読み取りに失敗しました。読めたふりはしません。（破損・非対応形式の可能性があります）",
                data={"path": path_text},
                error=str(error),
                executed=True,
            )

        truncated_text = len(text) > self.max_text_chars
        limited_text = text[: self.max_text_chars]

        if not limited_text.strip():
            return ToolResult(
                success=False,
                message=(
                    "PowerPointファイルは開けましたが、テキストを取り出せませんでした。"
                    "画像だけのスライドの可能性があります。"
                ),
                data={"path": path_text, "size_bytes": size_bytes, "slide_count": slide_count},
                error="empty_text",
                executed=True,
            )

        return ToolResult(
            success=True,
            message="PowerPointファイルを読み取りました。",
            data={
                "path": path_text,
                "resolved_path": str(path.resolve()),
                "file_name": path.name,
                "text": limited_text,
                "text_preview": limited_text[:500],
                "size_bytes": size_bytes,
                "slide_count": slide_count,
                "truncated": truncated_text,
            },
            error=None,
            executed=True,
        )

    def _extract_text(self, path):
        presentation = self._build_presentation_reader(path)
        slides = list(presentation.slides)[: self.max_slides]
        slide_texts = []

        for slide_index, slide in enumerate(slides, start=1):
            texts = []

            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = shape.text_frame.text

                    if text and text.strip():
                        texts.append(text.strip())

                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        row_text = " | ".join(
                            cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()
                        )

                        if row_text:
                            texts.append(row_text)

            notes_slide = getattr(slide, "notes_slide", None)

            if notes_slide is not None and getattr(notes_slide, "notes_text_frame", None) is not None:
                notes_text = notes_slide.notes_text_frame.text

                if notes_text and notes_text.strip():
                    texts.append(f"[ノート] {notes_text.strip()}")

            if texts:
                slide_texts.append(f"[スライド{slide_index}]\n" + "\n".join(texts))

        return len(slides), "\n\n".join(slide_texts)

    def _build_presentation_reader(self, path):
        if self._presentation_reader_factory is not None:
            return self._presentation_reader_factory(path)

        from pptx import Presentation

        return Presentation(str(path))

    def _get_path_from_request(self, request):
        input_data = getattr(request, "input_data", {}) or {}
        path_text = input_data.get("path")

        if path_text:
            return str(path_text).strip()

        original_input = input_data.get("original_input", "")
        return extract_pptx_path_from_text(original_input)
