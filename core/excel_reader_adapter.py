from pathlib import Path

from core.file_path_extraction import extract_single_file_path_from_text
from core.tool_adapter import ToolAdapter, ToolResult


EXCEL_EXTENSIONS = (".xlsx",)


def extract_excel_path_from_text(text):
    """テキストから、ユーザーが指定した単一のExcel(.xlsx)ファイルパスを取り出します。

    ディレクトリ全体やワイルドカードは対象にせず、拡張子で終わる1つのパスだけを拾います。
    スペースを含むファイル名（例: Windowsの重複ダウンロードによる「 (1).xlsx」）にも対応します。
    """
    return extract_single_file_path_from_text(text, EXCEL_EXTENSIONS)


class ExcelReaderAdapter(ToolAdapter):
    """最小実用版のExcel Reader Adapterです。

    ユーザーが明示した単一の.xlsxファイルだけを、確認後にテキスト抽出します。
    ディレクトリ探索、ワイルドカード、書き込み、実行は一切行いません。
    数式は評価値ではなく、保存されている計算済みの値を読み取ります（マクロは実行しません）。
    古い形式の.xls（バイナリ）には対応しません。
    """

    def __init__(
        self,
        max_file_bytes=20 * 1024 * 1024,
        max_rows_per_sheet=200,
        max_text_chars=12000,
        workbook_reader_factory=None,
    ):
        super().__init__(
            name="excel_reader",
            description="Read a single local Excel (.xlsx) file specified by the user.",
        )
        self.max_file_bytes = max_file_bytes
        self.max_rows_per_sheet = max_rows_per_sheet
        self.max_text_chars = max_text_chars
        # 自動checkでfakeのワークブック読み取りに差し替えられるようにします。
        self._workbook_reader_factory = workbook_reader_factory

    def is_available(self):
        return True

    def execute(self, request):
        """ToolRequestを受け取り、Excel本文をToolResultとして返します。"""
        path_text = self._get_path_from_request(request)

        if not path_text:
            return ToolResult(
                success=False,
                message="読み取るExcelファイルパスが見つからないため、Excel Readerは実行していません。",
                data=None,
                error="missing_path",
                executed=False,
            )

        path = Path(path_text)

        if path.suffix.lower() != ".xlsx":
            return ToolResult(
                success=False,
                message=(
                    "対応していない拡張子のため、Excel Readerは実行していません。"
                    "（対応拡張子: .xlsx。古い.xls形式には対応していません）"
                ),
                data={"path": path_text},
                error="unsupported_extension",
                executed=False,
            )

        if not path.exists():
            return ToolResult(
                success=False,
                message="指定されたファイルが見つからないため、Excel Readerは実行していません。",
                data={"path": path_text},
                error="file_not_found",
                executed=False,
            )

        if not path.is_file():
            return ToolResult(
                success=False,
                message="指定されたパスはファイルではないため、Excel Readerは実行していません。",
                data={"path": path_text},
                error="not_a_file",
                executed=False,
            )

        try:
            size_bytes = path.stat().st_size
        except OSError as error:
            return ToolResult(
                success=False,
                message="ファイル情報の取得に失敗しました。読めたふりはしません。",
                data={"path": path_text},
                error=str(error),
                executed=True,
            )

        if size_bytes > self.max_file_bytes:
            return ToolResult(
                success=False,
                message=(
                    "ファイルサイズが上限を超えているため、Excel Readerは実行していません。"
                    f"（上限: {self.max_file_bytes}バイト）"
                ),
                data={"path": path_text, "size_bytes": size_bytes},
                error="file_too_large",
                executed=False,
            )

        try:
            sheet_count, text = self._extract_text(path)
        except Exception as error:
            return ToolResult(
                success=False,
                message="Excelファイルの読み取りに失敗しました。読めたふりはしません。（破損・非対応形式の可能性があります）",
                data={"path": path_text},
                error=str(error),
                executed=True,
            )

        truncated_text = len(text) > self.max_text_chars
        limited_text = text[: self.max_text_chars]

        if not limited_text.strip():
            return ToolResult(
                success=False,
                message="Excelファイルは開けましたが、テキストを取り出せませんでした（空のシートの可能性があります）。",
                data={"path": path_text, "size_bytes": size_bytes, "sheet_count": sheet_count},
                error="empty_text",
                executed=True,
            )

        return ToolResult(
            success=True,
            message="Excelファイルを読み取りました。",
            data={
                "path": path_text,
                "resolved_path": str(path.resolve()),
                "file_name": path.name,
                "text": limited_text,
                "text_preview": limited_text[:500],
                "size_bytes": size_bytes,
                "sheet_count": sheet_count,
                "truncated": truncated_text,
            },
            error=None,
            executed=True,
        )

    def _extract_text(self, path):
        workbook = self._build_workbook_reader(path)

        try:
            sheet_names = getattr(workbook, "sheetnames", [])
            sheet_texts = []

            for sheet_name in sheet_names:
                sheet = workbook[sheet_name]
                row_texts = []

                for row in list(sheet.iter_rows(values_only=True))[: self.max_rows_per_sheet]:
                    cell_values = [str(value) for value in row if value is not None]

                    if cell_values:
                        row_texts.append(" | ".join(cell_values))

                if row_texts:
                    sheet_texts.append(f"[{sheet_name}]\n" + "\n".join(row_texts))

            return len(sheet_names), "\n\n".join(sheet_texts)
        finally:
            # Windowsではread_onlyワークブックのファイルハンドルが残りやすいため、明示的に閉じます。
            close = getattr(workbook, "close", None)

            if callable(close):
                close()

    def _build_workbook_reader(self, path):
        if self._workbook_reader_factory is not None:
            return self._workbook_reader_factory(path)

        from openpyxl import load_workbook

        return load_workbook(str(path), read_only=True, data_only=True)

    def _get_path_from_request(self, request):
        input_data = getattr(request, "input_data", {}) or {}
        path_text = input_data.get("path")

        if path_text:
            return str(path_text).strip()

        original_input = input_data.get("original_input", "")
        return extract_excel_path_from_text(original_input)
